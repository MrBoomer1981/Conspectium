"""Word extraction — app/services/docx.py.

No model is involved, so the whole contract is structural: which paragraph
styles become markdown headings, what is dropped, and how the title is
derived. Documents are built in memory with python-docx rather than kept as
binary fixtures, so the expectations stay readable.
"""

import io

import docx as python_docx

from app.services import docx


def _build(*paragraphs: tuple[str, str | None]) -> bytes:
    """A .docx from (text, style) pairs. style=None means body text."""
    document = python_docx.Document()
    for text, style in paragraphs:
        document.add_paragraph(text, style=style) if style else (
            document.add_paragraph(text)
        )
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_body_text_is_copied_unchanged():
    result = docx.extract_markdown(_build(("просто абзац", None)))
    assert result == "просто абзац"


def test_title_style_becomes_a_single_hash():
    result = docx.extract_markdown(_build(("Заголовок", "Title")))
    assert result == "# Заголовок"


def test_heading_levels_become_that_many_hashes():
    result = docx.extract_markdown(
        _build(("один", "Heading 1"), ("два", "Heading 2"), ("три", "Heading 3"))
    )
    assert result == "# один\n\n## два\n\n### три"


def test_heading_depth_is_capped_at_six():
    result = docx.extract_markdown(_build(("глубоко", "Heading 9")))
    assert result == "###### глубоко"


def test_blank_paragraphs_are_dropped():
    result = docx.extract_markdown(
        _build(("первый", None), ("   ", None), ("", None), ("второй", None))
    )
    assert result == "первый\n\nвторой"


def test_paragraphs_are_separated_by_a_blank_line():
    result = docx.extract_markdown(_build(("а", None), ("б", None)))
    assert result == "а\n\nб"


def test_empty_document_yields_empty_string():
    assert docx.extract_markdown(_build()) == ""


def test_whitespace_around_text_is_stripped():
    assert docx.extract_markdown(_build(("  с пробелами  ", None))) == "с пробелами"


# ── derive_title ─────────────────────────────────────────────────────────────


def test_title_is_the_first_non_empty_line():
    assert docx.derive_title("первая\n\nвторая", "запасной") == "первая"


def test_title_strips_heading_marks():
    assert docx.derive_title("### Тема урока", "запасной") == "Тема урока"


def test_title_skips_leading_blank_lines():
    assert docx.derive_title("\n\n  \nНастоящая строка", "з") == "Настоящая строка"


def test_title_is_capped_at_60_characters():
    assert len(docx.derive_title("я" * 200, "з")) == 60


def test_title_falls_back_when_markdown_is_blank():
    assert docx.derive_title("   \n\n  ", "имя файла") == "имя файла"


def test_title_fallback_is_also_capped():
    assert len(docx.derive_title("", "ф" * 200)) == 60


def test_title_is_untitled_when_there_is_nothing_at_all():
    assert docx.derive_title("", "") == "Untitled"
